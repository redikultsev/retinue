"""What the owner sends besides text, made readable for the model — by code, in the router, before any model run.

Photos and PDFs reach the model as content blocks; office and text files as their text; voice, audio and the sound
of a video as a transcript. Whatever cannot be read is refused aloud, with the reason. The model never gets a file
it would have to open itself: the assistant has no file tools.

Pictures, PDFs and office files are parsed by untrusted-input libraries (libheif, pypdf, lxml) in a separate process
(`python -m retinue.attachments`) with a deadline and a memory ceiling: a file that hangs, swells or crashes the
parser costs that process and is refused aloud; the router goes on.
"""

from __future__ import annotations

import asyncio
import base64
import io
import json
import os
import sys
import zipfile
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import PurePosixPath

from PIL import Image, ImageOps
from pillow_heif import register_heif_opener
from pypdf import PdfReader, PdfWriter

from .speech import SpeechError

register_heif_opener()  # Pillow opens HEIC/HEIF from here on: an iPhone photo sent as a file

MAX_DOWNLOAD = 20 * 1024 * 1024  # the Bot API serves bots files up to 20 MB (getFile)
IMAGE_SIDE = 1568        # the long side of a picture the model gets: ≤ 2352 tokens for 4:3 (research 54, §3)
JPEG_QUALITY = 85
MAX_PIXELS = 50_000_000  # a bigger picture is refused before it is decoded: a 20 MB PNG can unpack to gigabytes
PDF_PAGES = 20           # a longer PDF goes as its text: a page costs 1.5–3 thousand tokens on every turn
PDF_BYTES = 5 * 1024 * 1024
PDF_SCAN_PAGES = 300     # a PDF's text layer is read from this many pages at most
# Text taken from one file: a document, a PDF's text layer, a transcript. Below the 50 000 characters up to which
# Claude Code keeps a tool result inline, so get_attachment returns it whole.
TEXT_CHARS = 30_000
UNZIPPED_BYTES = 50 * 1024 * 1024  # docx, xlsx and pptx are zip archives: a bigger inside is a zip bomb
XLSX_ROWS = 20_000       # rows of a workbook read at most, empty ones included
PARSE_TIMEOUT_S = 60     # a file the worker has not read by then is refused
PARSE_MEMORY = 1536 * 1024 * 1024  # the worker's address space; glibc arenas are capped too (MALLOC_ARENA_MAX)
WORKER = [sys.executable, "-m", "retinue.attachments"]
OFFICE = {".docx": "документ", ".xlsx": "таблица", ".pptx": "презентация"}
TEXT = {".txt", ".md", ".csv", ".tsv", ".json", ".html", ".htm"}
PICTURES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic"}
# What is sent to speech to text, as (file name, media type, how the mark names it). Telegram calls a voice note
# .oga; ElevenLabs may judge the format by the name, so the name says .ogg. Audio keeps its own name and type.
SPEECH = {"voice": ("voice.ogg", "audio/ogg", "голосовое"), "video_note": ("video.mp4", "video/mp4", "видеосообщение"),
          "video": ("video.mp4", "video/mp4", "видео"), "audio": ("", "", "аудио")}
KINDS = {"photo": "фото", "document": "файл", **{kind: label for kind, (_, _, label) in SPEECH.items()}}
TEXT_TYPES = {"text/plain": ".txt", "text/markdown": ".md", "text/csv": ".csv", "text/tab-separated-values": ".tsv",
              "application/json": ".json", "text/html": ".html"}


class Unreadable(Exception):
    """The file cannot be given to the model. The text finishes a sentence to the owner."""


@dataclass
class Prepared:
    """One file, ready for the model."""

    what: str                                   # how it is named in the mark: «фото 1280×960», «PDF «a.pdf», 3 стр.»
    files: list[tuple[str, bytes]] = field(default_factory=list)  # (media type, bytes): content blocks — JPEG, PDF
    text: str = ""                              # what the model reads as text: a transcript, a document's text


@dataclass
class Upload:
    """One file as a channel received it."""

    kind: str                # photo | document | voice | audio | video | video_note
    data: bytes = b""
    name: str = ""           # the file's own name, when the messenger has one
    media_type: str = ""
    duration: int = 0        # seconds, as the messenger says
    size: int = 0            # bytes, as the messenger announced them before the download
    refused: str = ""        # the channel could not fetch it: why, finishing «Не прочитано: …»
    # Downloads the bytes when the core needs them: an album is gathered first, the parts are fetched after.
    fetch: Callable[[], Awaitable[bytes]] | None = field(default=None, compare=False, repr=False)

    def label(self) -> str:
        """«видео «clip.mp4»», «голосовое», «фото»"""
        kind = KINDS.get(self.kind, self.kind)
        return f"{kind} «{self.name}»" if self.name and self.kind != "photo" else kind


def too_big(upload: Upload) -> str:
    return f"{upload.label()}, {round(upload.size / 2**20)} МБ: Telegram отдаёт ботам файлы до {MAX_DOWNLOAD // 2**20} МБ"


def mmss(seconds: int) -> str:
    """0:42, 3:10, 1:02:05"""
    minutes, seconds = divmod(int(seconds), 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours}:{minutes:02d}:{seconds:02d}" if hours else f"{minutes}:{seconds:02d}"


async def prepare(upload: Upload, scribe) -> Prepared:
    """One upload -> what the model gets. Raises Unreadable with the reason, which the router tells the owner.
    `scribe` turns speech into text (`speech.Scribe`); without it nothing is transcribed."""
    if upload.refused:
        raise Unreadable(upload.refused)
    if upload.fetch and not upload.data:
        try:
            upload.data = await upload.fetch()
        except Exception as exc:  # the channel words its errors without addresses: they carry tokens
            raise Unreadable(f"{upload.label()}: Telegram не отдал его ({exc})") from exc
    if upload.kind in ("photo", "document"):
        return await isolated(upload)
    if upload.kind in SPEECH:
        filename, media_type, _ = SPEECH[upload.kind]
        what = f"{upload.label()} {mmss(upload.duration)}"
        try:
            if scribe is None:
                raise SpeechError("расшифровка не настроена")
            text = await scribe.transcribe(upload.data, filename or upload.name or "audio",
                                           media_type or upload.media_type or "application/octet-stream")
        except SpeechError as exc:
            raise Unreadable(f"{what}: {exc}") from exc
        return Prepared(what, text=cut(text) or "(речи не разобрать)")
    raise Unreadable(f"не умею читать такие сообщения: {upload.kind}")


def read(upload: Upload) -> Prepared:
    """A picture, a PDF or a document -> what the model gets. The worker runs this; raises Unreadable."""
    suffix = PurePosixPath(upload.name.lower()).suffix
    if upload.kind == "photo" or upload.media_type.startswith("image/") or suffix in PICTURES:
        label = "фото" if upload.kind == "photo" else f"картинка «{upload.name}»"
        try:
            data, width, height = image(upload.data)
        except Unreadable as exc:
            raise Unreadable(f"{label}: {exc}") from exc
        return Prepared(f"{label}{' ' if upload.kind == 'photo' else ', '}{width}×{height}", [("image/jpeg", data)])
    if upload.media_type == "application/pdf" or suffix == ".pdf":
        return pdf(upload.data, upload.name or "без имени")
    return document(upload.data, upload.name or "без имени", upload.media_type)


def _parsing(upload: Upload) -> str:
    """How a file is named in a refusal about its parsing: «фото», «картинка «a.heic»», «файл «a.pdf»»."""
    suffix = PurePosixPath(upload.name.lower()).suffix
    if upload.kind == "document" and (upload.media_type.startswith("image/") or suffix in PICTURES):
        return f"картинка «{upload.name}»"
    return upload.label()


async def isolated(upload: Upload) -> Prepared:
    """`read` in the worker process: the bytes go in through stdin, the result comes back as JSON (never pickle:
    the worker reads hostile files). Too slow, crashed or out of memory -> Unreadable with the reason."""
    request = json.dumps({"kind": upload.kind, "name": upload.name, "media_type": upload.media_type,
                          "data": base64.b64encode(upload.data).decode()}).encode()
    process = await asyncio.create_subprocess_exec(
        *WORKER, stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.DEVNULL,
        env={**os.environ, "MALLOC_ARENA_MAX": "2"})
    try:
        out, _ = await asyncio.wait_for(process.communicate(request), PARSE_TIMEOUT_S)
    except TimeoutError:
        process.kill()
        await process.wait()
        raise Unreadable(f"{_parsing(upload)}: разбор не уложился в {PARSE_TIMEOUT_S} с") from None
    try:
        answer = json.loads(out)
    except ValueError:
        raise Unreadable(f"{_parsing(upload)}: разбор упал (код {process.returncode})") from None
    if not answer.get("ok"):
        raise Unreadable(answer.get("reason") or f"{_parsing(upload)}: разбор не удался")
    return Prepared(answer["what"], [(media_type, base64.b64decode(data)) for media_type, data in answer["files"]],
                    answer["text"])


def worker() -> None:
    """`python -m retinue.attachments`: one request on stdin, one JSON answer on stdout."""
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (PARSE_MEMORY, PARSE_MEMORY))
    except (ImportError, ValueError, OSError):
        pass  # macOS does not limit the address space; the deadline still holds
    request = json.loads(sys.stdin.buffer.read())
    upload = Upload(request["kind"], base64.b64decode(request["data"]), request["name"], request["media_type"])
    try:
        done = read(upload)
        answer = {"ok": True, "what": done.what, "text": done.text,
                  "files": [[media_type, base64.b64encode(data).decode()] for media_type, data in done.files]}
    except Unreadable as exc:
        answer = {"ok": False, "reason": str(exc)}
    except MemoryError:
        answer = {"ok": False, "reason": f"{_parsing(upload)}: на разбор не хватило памяти"}
    except Exception as exc:
        answer = {"ok": False, "reason": f"{_parsing(upload)}: ошибка разбора ({type(exc).__name__})"}
    sys.stdout.write(json.dumps(answer))


def cut(text: str) -> str:
    text = text.strip()
    if len(text) <= TEXT_CHARS:
        return text
    return f"{text[:TEXT_CHARS]}\n…(обрезано на {TEXT_CHARS} знаках)"


def image(data: bytes) -> tuple[bytes, int, int]:
    """Any picture Pillow opens -> (JPEG, width, height): turned upright by its EXIF (the model does not read
    EXIF), the long side at most IMAGE_SIDE, transparency on white. Re-encoding is also the check: a picture that
    does not decode here never reaches the session, where a refused block would break every later turn."""
    try:
        with Image.open(io.BytesIO(data)) as img:
            if img.width * img.height > MAX_PIXELS:
                raise Unreadable(f"картинка слишком большая: {img.width}×{img.height}")
            img = ImageOps.exif_transpose(img)
            img.thumbnail((IMAGE_SIDE, IMAGE_SIDE))
            if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
                rgba = img.convert("RGBA")
                flat = Image.new("RGB", rgba.size, (255, 255, 255))
                flat.paste(rgba, mask=rgba.getchannel("A"))
                img = flat
            elif img.mode != "RGB":
                img = img.convert("RGB")
            out = io.BytesIO()
            img.save(out, "JPEG", quality=JPEG_QUALITY)
            return out.getvalue(), img.width, img.height
    except Unreadable:
        raise
    except Exception as exc:  # not a picture, a broken one, a decompression bomb
        raise Unreadable("картинку не удалось открыть") from exc


def pdf(data: bytes, name: str) -> Prepared:
    """A short PDF goes whole: the model reads its pages, scans included. A long one goes as its text layer; a long
    scan has none and is refused. The text layer is kept either way: get_attachment returns it after compaction,
    when the document block is gone from the session."""
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise Unreadable(f"PDF «{name}» под паролем")
        pages = len(reader.pages)
        layer, chars = [], 0
        for page in reader.pages[:PDF_SCAN_PAGES]:
            if chars > TEXT_CHARS:
                break  # a thousand-page book is not read to the end for a text that is cut anyway
            layer.append(page.extract_text() or "")
            chars += len(layer[-1])
        whole = pages <= PDF_PAGES and len(data) <= PDF_BYTES
        if whole and reader.is_encrypted:  # opens without a password, only editing is locked: a plain copy
            out = io.BytesIO()
            PdfWriter(clone_from=reader).write(out)
            data = out.getvalue()
    except Unreadable:
        raise
    except Exception as exc:
        raise Unreadable(f"PDF «{name}» не открывается: файл повреждён") from exc
    text = cut("\n".join(layer))
    if len(layer) < pages and chars <= TEXT_CHARS and text:
        text += f"\n…(текст только первых {len(layer)} стр. из {pages})"
    label = f"PDF «{name}», {pages} стр."
    if whole:
        return Prepared(label, [("application/pdf", data)], text)
    if text:
        return Prepared(f"{label}, только текст", text=text)
    raise Unreadable(f"PDF «{name}», {pages} стр.: больше {PDF_PAGES} стр. или {PDF_BYTES // 2**20} МБ, а текста в нём "
                     "нет — это скан. Пришли нужные страницы фотографиями")


def document(data: bytes, name: str, media_type: str) -> Prepared:
    """An office or text file -> its text. The type is taken from the name, then from the media type; any other
    type is refused with both named."""
    suffix = PurePosixPath(name.lower()).suffix
    if suffix not in OFFICE and suffix not in TEXT:
        suffix = TEXT_TYPES.get(media_type.split(";")[0].strip(), suffix)
    if suffix in OFFICE:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                inside = sum(item.file_size for item in archive.infolist())
            if inside > UNZIPPED_BYTES:
                raise Unreadable(f"«{name}» распаковывается больше чем в {UNZIPPED_BYTES // 2**20} МБ")
            text = {".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx}[suffix](data)
        except Unreadable:
            raise
        except Exception as exc:
            raise Unreadable(f"«{name}» не открывается: файл повреждён или это не {suffix}") from exc
        return Prepared(f"{OFFICE[suffix]} «{name}»", text=cut(text))
    if suffix in TEXT:
        text = _decode(data)
        return Prepared(f"файл «{name}»", text=cut(_html(text) if suffix in (".html", ".htm") else text))
    raise Unreadable(f"не умею читать такие файлы: «{name}» ({media_type or 'тип не назван'})")


def _decode(data: bytes) -> str:
    """UTF-8 (a BOM is dropped), else Windows-1251: what a Russian Excel or Notepad writes."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


class _Text(HTMLParser):
    BLOCKS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        self.skip += tag in ("script", "style")
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        self.skip -= tag in ("script", "style") and self.skip > 0
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def _html(text: str) -> str:
    parser = _Text()
    parser.feed(text)
    lines = (" ".join(line.split(" ")).strip() for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def _docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table

    lines = []
    for block in Document(io.BytesIO(data)).iter_inner_content():
        if isinstance(block, Table):
            lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
        elif block.text.strip():
            lines.append(block.text.strip())
    return "\n".join(lines)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime) and not (value.hour or value.minute or value.second):
        return value.date().isoformat()
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # data_only: values, not formulas
    lines, chars, rows = [], 0, 0
    for sheet in book.worksheets:
        sheet.reset_dimensions()  # the size a sheet claims is not trusted: «A1:XFD1048576» pads every row
        lines.append(f"## Лист «{sheet.title}»")
        for row in sheet.iter_rows(values_only=True):
            if chars > TEXT_CHARS:
                break
            if rows >= XLSX_ROWS:
                lines.append(f"…(дальше строки не читались: больше {XLSX_ROWS})")
                book.close()
                return "\n".join(lines)
            rows += 1
            cells = list(row)
            while cells and cells[-1] is None:
                cells.pop()
            if cells:
                lines.append(" | ".join(_cell(value) for value in cells))
                chars += len(lines[-1])
    book.close()
    return "\n".join(lines)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    lines = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        lines.append(f"## Слайд {number}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append(f"Заметки: {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(lines)


if __name__ == "__main__":
    worker()
