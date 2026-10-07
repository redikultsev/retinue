"""Attachments without the network: what the model gets from a photo, a PDF, a document, and what is refused."""

import asyncio
import io
import sys
from datetime import date

import pytest
from PIL import Image
from pillow_heif import register_heif_opener
from pypdf import PdfReader, PdfWriter

from retinue import attachments
from retinue.attachments import Upload, Unreadable, prepare, too_big
from retinue.speech import SpeechError

register_heif_opener()  # so that the test can write HEIF too


def picture(w: int, h: int, fmt: str = "JPEG", mode: str = "RGB", orientation: int | None = None) -> bytes:
    img = Image.new(mode, (w, h), (200, 30, 30, 0) if mode == "RGBA" else (200, 30, 30))
    exif = Image.Exif()
    if orientation:
        exif[0x0112] = orientation  # how the camera was held
    out = io.BytesIO()
    img.save(out, fmt, exif=exif)
    return out.getvalue()


def test_a_photo_is_turned_upright_shrunk_and_made_jpeg():
    data, w, h = attachments.image(picture(4000, 3000, orientation=6))  # taken with the phone on its side
    assert (w, h) == (1176, 1568), "the rotation is applied, then the long side is cut to 1568"
    with Image.open(io.BytesIO(data)) as out:
        assert out.format == "JPEG" and out.size == (1176, 1568) and 0x0112 not in out.getexif()
    assert attachments.image(picture(800, 600))[1:] == (800, 600), "a small photo is not enlarged"


def test_any_picture_becomes_jpeg_and_garbage_is_refused(monkeypatch):
    data, w, h = attachments.image(picture(100, 50, "PNG", "RGBA"))
    with Image.open(io.BytesIO(data)) as out:
        assert (out.format, out.mode, out.size) == ("JPEG", "RGB", (100, 50))
        assert out.getpixel((0, 0)) == (255, 255, 255), "transparency becomes white, not black"
    assert attachments.image(picture(64, 64, "WEBP"))[1:] == (64, 64)
    heic, w, h = attachments.image(picture(300, 200, "HEIF"))  # an iPhone photo sent as a file
    assert (w, h) == (300, 200) and Image.open(io.BytesIO(heic)).format == "JPEG"
    with pytest.raises(Unreadable, match="картинку не удалось открыть"):
        attachments.image(b"GIF89a not really")
    monkeypatch.setattr(attachments, "MAX_PIXELS", 100 * 100)
    with pytest.raises(Unreadable, match="слишком большая: 200×200"):
        attachments.image(picture(200, 200, "PNG"))


def make_pdf(pages: list[str]) -> bytes:
    """A real PDF with a text layer: one line of Helvetica per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>",
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET"
        objects += [f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >>"
                    f" /Contents {5 + 2 * i} 0 R >>", f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    table = "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    return out + (f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n{table}"
                  f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{len(out)}\n%%EOF\n").encode()


def scan(pages: int) -> bytes:
    """A PDF of pictures and no text, as a phone scanner makes it."""
    images = [Image.new("RGB", (200, 280), (240, 240, 240)) for _ in range(pages)]
    out = io.BytesIO()
    images[0].save(out, "PDF", save_all=True, append_images=images[1:])
    return out.getvalue()


def test_a_small_pdf_goes_to_the_model_whole():
    data = make_pdf(["Invoice 4711", "Total 99 EUR"])
    ready = attachments.pdf(data, "invoice.pdf")
    assert ready.what == "PDF «invoice.pdf», 2 стр." and ready.files == [("application/pdf", data)]
    assert "Invoice 4711" in ready.text and "Total 99 EUR" in ready.text, "the text layer is kept: for get_attachment"
    assert attachments.pdf(scan(2), "scan.pdf").files, "a short scan goes whole too: the model reads the pictures"


def test_a_long_pdf_is_read_as_text_and_a_long_scan_is_refused(monkeypatch):
    monkeypatch.setattr(attachments, "PDF_PAGES", 2)
    long = attachments.pdf(make_pdf(["one", "two", "three"]), "book.pdf")
    assert long.what == "PDF «book.pdf», 3 стр., только текст" and long.files == [] and "three" in long.text
    with pytest.raises(Unreadable, match="«scan.pdf», 3 стр.: больше 2 стр. или 5 МБ, а текста в нём нет"):
        attachments.pdf(scan(3), "scan.pdf")
    monkeypatch.setattr(attachments, "TEXT_CHARS", 20)
    cut = attachments.pdf(make_pdf(["0123456789ABCDEF"] * 3), "long.pdf").text
    assert cut.startswith("0123456789ABCDEF") and cut.endswith("\n…(обрезано на 20 знаках)")


def test_a_locked_or_broken_pdf():
    def locked(user_password: str) -> bytes:
        writer = PdfWriter()
        writer.add_blank_page(200, 200)
        writer.encrypt(user_password=user_password, owner_password="owner", algorithm="RC4-128")
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue()

    with pytest.raises(Unreadable, match="PDF «a.pdf» под паролем"):
        attachments.pdf(locked("secret"), "a.pdf")
    # A bank statement: anyone may open it, only editing is locked. The model gets an unlocked copy.
    ((media_type, copy),) = attachments.pdf(locked(""), "statement.pdf").files
    assert media_type == "application/pdf" and not PdfReader(io.BytesIO(copy)).is_encrypted
    with pytest.raises(Unreadable, match="PDF «b.pdf» не открывается"):
        attachments.pdf(b"%PDF-1.4 not really", "b.pdf")


def office(kind: str) -> bytes:
    """A small document of each kind, made by the same libraries that read it."""
    out = io.BytesIO()
    if kind == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("Договор аренды № 17")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Сумма", "900 EUR"
        doc.add_paragraph("Подписано в Белграде")
        doc.save(out)
    elif kind == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Расходы"
        book.active.append(["Дата", "Сумма"])
        book.active.append([date(2026, 10, 1), 1250.5])
        book.create_sheet("Пустой")
        book.save(out)
    else:
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "План на квартал"
        slide.placeholders[1].text = "Запустить почту"
        slide.notes_slide.notes_text_frame.text = "сказать про сроки"
        deck.save(out)
    return out.getvalue()


def test_office_files_are_read_as_text():
    docx = attachments.document(office("docx"), "договор.docx", "")
    assert docx.what == "документ «договор.docx»" and docx.files == []
    assert docx.text == "Договор аренды № 17\nСумма | 900 EUR\nПодписано в Белграде", "tables in their place"
    xlsx = attachments.document(office("xlsx"), "расходы.xlsx", "")
    assert xlsx.what == "таблица «расходы.xlsx»"
    assert xlsx.text == "## Лист «Расходы»\nДата | Сумма\n2026-10-01 | 1250.5\n## Лист «Пустой»"
    pptx = attachments.document(office("pptx"), "план.pptx", "")
    assert pptx.what == "презентация «план.pptx»"
    assert pptx.text == "## Слайд 1\nПлан на квартал\nЗапустить почту\nЗаметки: сказать про сроки"


def test_text_files_in_any_usual_encoding():
    csv = attachments.document("имя;сумма\nАня;10\n".encode("cp1251"), "список.csv", "text/csv")
    assert csv.what == "файл «список.csv»" and csv.text == "имя;сумма\nАня;10", "a Russian Excel export is cp1251"
    assert attachments.document("\ufeffзаметка".encode(), "a.txt", "").text == "заметка", "the BOM is dropped"
    page = "<html><head><style>p{}</style><script>alert(1)</script></head><body><h1>Билет</h1><p>Рейс&nbsp;JU 680</p></body></html>"
    assert attachments.document(page.encode(), "билет.html", "text/html").text == "Билет\nРейс\xa0JU 680"
    assert attachments.document(b'{"a": 1}', "data", "application/json").text == '{"a": 1}', "by the type, too"


def test_unknown_broken_and_swollen_files_are_refused(monkeypatch):
    with pytest.raises(Unreadable, match="не умею читать такие файлы: «setup.exe» \\(application/x-msdownload\\)"):
        attachments.document(b"MZ", "setup.exe", "application/x-msdownload")
    with pytest.raises(Unreadable, match="«отчёт.docx» не открывается"):
        attachments.document(b"PK not a zip", "отчёт.docx", "")
    monkeypatch.setattr(attachments, "UNZIPPED_BYTES", 1000)
    with pytest.raises(Unreadable, match="«договор.docx» распаковывается больше чем в 0 МБ"):
        attachments.document(office("docx"), "договор.docx", "")


class FakeScribe:
    """Speech to text that remembers what it was given."""

    def __init__(self, text: str = "", error: str = ""):
        self.text, self.error, self.calls = text, error, []

    async def transcribe(self, data, filename, media_type):
        self.calls.append((filename, media_type, data))
        if self.error:
            raise SpeechError(self.error)
        return self.text


def ready(upload, scribe=None):
    return asyncio.run(prepare(upload, scribe))


def test_each_kind_of_upload_becomes_blocks_or_text():
    photo = ready(Upload("photo", picture(1280, 960)))
    assert photo.what == "фото 1280×960" and [t for t, _ in photo.files] == ["image/jpeg"] and photo.text == ""
    drawing = ready(Upload("document", picture(3000, 2000, "PNG"), "схема.png", "image/png"))
    assert drawing.what == "картинка «схема.png», 1568×1045" and drawing.files[0][0] == "image/jpeg", "a file, then a photo"
    paper = ready(Upload("document", make_pdf(["Invoice 4711"]), "invoice", "application/pdf"))
    assert paper.what == "PDF «invoice», 1 стр." and paper.files[0][0] == "application/pdf"
    assert ready(Upload("document", office("docx"), "договор.docx")).text.startswith("Договор аренды")

    scribe = FakeScribe(" купи хлеб ")
    voice = ready(Upload("voice", b"OggS", media_type="audio/ogg", duration=42), scribe)
    assert (voice.what, voice.text, voice.files) == ("голосовое 0:42", "купи хлеб", [])
    ready(Upload("video_note", b"mp4", duration=15), scribe)
    ready(Upload("audio", b"ID3", "lecture.mp3", "audio/mpeg", duration=190), scribe)
    assert [c[:2] for c in scribe.calls] == [("voice.ogg", "audio/ogg"), ("video.mp4", "video/mp4"),
                                             ("lecture.mp3", "audio/mpeg")], "Telegram's .oga is sent as voice.ogg"
    assert ready(Upload("audio", b"ID3", "lecture.mp3", "audio/mpeg", duration=190), scribe).what == "аудио «lecture.mp3» 3:10"
    assert ready(Upload("voice", b"OggS", duration=3), FakeScribe("")).text == "(речи не разобрать)"


def test_what_cannot_be_read_says_why():
    def why(upload, scribe=None):
        with pytest.raises(Unreadable) as refused:
            ready(upload, scribe)
        return str(refused.value)

    huge = Upload("video", name="clip.mp4", size=35 * 2**20)
    huge.refused = too_big(huge)
    assert why(huge) == "видео «clip.mp4», 35 МБ: Telegram отдаёт ботам файлы до 20 МБ"
    assert why(Upload("voice", b"OggS", duration=42), FakeScribe(error="кончились кредиты ElevenLabs")) == \
        "голосовое 0:42: кончились кредиты ElevenLabs"
    assert why(Upload("voice", b"OggS", duration=5)) == "голосовое 0:05: расшифровка не настроена"
    assert why(Upload("photo", b"broken")) == "фото: картинку не удалось открыть"
    assert why(Upload("document", b"MZ", "setup.exe", "application/x-msdownload")).startswith("не умею читать такие файлы")


# --- review 2026-10-07 --------------------------------------------------------------------------------------------


def stretched_xlsx(rows: int) -> bytes:
    """One cell at the bottom of a sheet that claims to be A1:XFD<rows>: openpyxl pads every row to 16 384 columns."""
    import re
    import zipfile

    from openpyxl import Workbook

    book, out = Workbook(), io.BytesIO()
    book.active[f"A{rows}"] = "y"
    book.save(out)
    source, swollen = zipfile.ZipFile(io.BytesIO(out.getvalue())), io.BytesIO()
    with zipfile.ZipFile(swollen, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', f'<dimension ref="A1:XFD{rows}"'.encode(), data)
            target.writestr(item, data)
    return swollen.getvalue()


def test_a_sheet_that_lies_about_its_size_is_read_by_its_cells(monkeypatch):
    assert attachments.document(stretched_xlsx(3000), "a.xlsx", "").text == "## Лист «Sheet»\ny", \
        "the claimed width is ignored, empty cells at the end of a row are dropped"
    monkeypatch.setattr(attachments, "XLSX_ROWS", 2)
    many = io.BytesIO()
    from openpyxl import Workbook
    book = Workbook()
    for number in range(5):
        book.active.append([number])
    book.save(many)
    assert attachments.document(many.getvalue(), "b.xlsx", "").text == \
        "## Лист «Sheet»\n0\n1\n…(дальше строки не читались: больше 2)"
    assert attachments.UNZIPPED_BYTES == 50 * 2**20


def test_a_long_pdf_is_read_only_so_far(monkeypatch):
    monkeypatch.setattr(attachments, "PDF_PAGES", 1)
    monkeypatch.setattr(attachments, "PDF_SCAN_PAGES", 2)
    text = attachments.pdf(make_pdf(["one", "two", "three"]), "book.pdf").text
    assert "three" not in text and text.endswith("…(текст только первых 2 стр. из 3)")


def test_an_aes_pdf_and_a_failing_unlock(monkeypatch):
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.encrypt(user_password="", owner_password="owner", algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    ((_, copy),) = attachments.pdf(out.getvalue(), "aes.pdf").files
    assert not PdfReader(io.BytesIO(copy)).is_encrypted, "AES needs `cryptography`: it is a declared dependency"

    class Broken:
        def __init__(self, **kwargs):
            raise ValueError("cannot clone")

    monkeypatch.setattr(attachments, "PdfWriter", Broken)
    with pytest.raises(Unreadable, match="PDF «aes.pdf» не открывается"):
        attachments.pdf(out.getvalue(), "aes.pdf")


def test_files_are_read_in_a_separate_process_with_a_deadline(monkeypatch):
    photo = ready(Upload("photo", picture(640, 480)))
    assert photo.what == "фото 640×480" and photo.files[0][0] == "image/jpeg", "the worker answers as before"
    monkeypatch.setattr(attachments, "PARSE_TIMEOUT_S", 0.5)
    monkeypatch.setattr(attachments, "WORKER", [sys.executable, "-c", "import time; time.sleep(30)"])
    with pytest.raises(Unreadable, match="фото: разбор не уложился в 0.5 с"):
        ready(Upload("photo", picture(10, 10)))
    monkeypatch.setattr(attachments, "WORKER", [sys.executable, "-c", "import os, signal; os.kill(os.getpid(), signal.SIGSEGV)"])
    with pytest.raises(Unreadable, match="картинка «a.heic»: разбор упал"):
        ready(Upload("document", b"x", "a.heic"))
